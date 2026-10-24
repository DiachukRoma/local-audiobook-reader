"""Plain-text and DOCX loaders."""

from __future__ import annotations

import re
from pathlib import Path

from .model import Document, DocumentError, Paragraph, Section
from .normalize import is_readable, normalize_text

_ENCODINGS = ("utf-8-sig", "utf-16", "cp1251", "koi8-u", "koi8-r")
_BLANK_LINE = re.compile(r"\n[ \t ]*\n")


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in _ENCODINGS:
        if encoding == "utf-16":
            continue
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding.startswith("koi8"):
            # KOI8 decodes any byte sequence; only accept it if the result looks Cyrillic.
            letters = [c for c in text if c.isalpha()]
            if letters and sum("а" <= c.lower() <= "я" for c in letters) / len(letters) < 0.5:
                continue
        return text
    return data.decode("utf-8", errors="replace")


def split_plain_paragraphs(text: str) -> list[str]:
    """Split raw text into paragraphs.

    Blank lines separate paragraphs. If the text has no blank lines at all, every
    non-empty line is treated as a paragraph (typical for many .txt books).
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = _BLANK_LINE.split(text)
    if len(blocks) == 1:
        blocks = text.split("\n")
    paragraphs = []
    for block in blocks:
        cleaned = normalize_text(block)
        if is_readable(cleaned):
            paragraphs.append(cleaned)
    return paragraphs


def load_txt(path: Path) -> Document:
    text = _decode(path.read_bytes())
    paragraphs = [Paragraph(p) for p in split_plain_paragraphs(text)]
    if not paragraphs:
        raise DocumentError("Файл не містить тексту.")
    return Document(path=path, title=path.stem, format="txt", paragraphs=paragraphs)


_HEADING_STYLES = ("heading", "title", "заголовок", "назва")


def load_docx(path: Path) -> Document:
    import docx  # python-docx

    try:
        doc = docx.Document(str(path))
    except Exception as exc:  # python-docx raises several exception types for broken files
        raise DocumentError(f"Не вдалося відкрити DOCX: {exc}") from exc

    paragraphs: list[Paragraph] = []
    sections: list[Section] = []
    for para in doc.paragraphs:
        text = normalize_text(para.text)
        if not text:
            continue
        style = (para.style.name if para.style is not None else "").lower()
        if style.startswith(_HEADING_STYLES):
            sections.append(Section(text, len(paragraphs)))
        paragraphs.append(Paragraph(text, section=len(sections) - 1))
    if not paragraphs:
        raise DocumentError("Документ не містить тексту.")

    title = (doc.core_properties.title or "").strip() or path.stem
    if paragraphs[0].section < 0:
        # Text before the first heading forms its own leading section.
        sections.insert(0, Section(title, 0))
        paragraphs = [Paragraph(p.text, p.section + 1) for p in paragraphs]
    return Document(path=path, title=title, format="docx", paragraphs=paragraphs, sections=sections)
